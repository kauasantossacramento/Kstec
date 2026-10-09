from datetime import timedelta
from decimal import Decimal

import pytest
from django.core import mail
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from apps.catalogo.importador import importar_arquivo
from apps.catalogo.models import (
    CodigoNBS,
    CodigoServicoLC116,
    CodigoTributacaoNacional,
    CorrelacaoNBS,
    ItemCatalogo,
)
from apps.catalogo.seeds import seed as seed_catalogo
from apps.catalogo.services import calcular_preco, sugestoes_fiscais
from apps.comercial import services
from apps.comercial.models import Orcamento


@pytest.fixture
def catalogo(empresa):
    seed_catalogo(empresa)
    return ItemCatalogo.objects.get(codigo_interno="GR-CARTAO")


def test_importador_anexos(db, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "LISTA.SERV.NAC."
    ws.append(["Código", "Descrição"])
    ws.append(["01", "Serviços de informática e congêneres"])
    ws.append(["01.07", "Suporte técnico em informática"])
    ws.append(["01.07.01", "Suporte técnico em informática, inclusive instalação"])
    ws.append(["13.05.01", "Composição gráfica personalizada"])
    nbs = wb.create_sheet("LISTA.NBS_v2.0")
    nbs.append(["NBS", "DESCRIÇÃO"])
    nbs.append(["1.15", "Serviços de TI (agregador)"])
    nbs.append(["1.1501.30.00", "Serviços de suporte técnico em TI"])
    arq_b = tmp_path / "ANEXO_B.xlsx"
    wb.save(arq_b)

    wb2 = Workbook()
    t = wb2.active
    t.title = "tabela geral"
    t.append(["Correlação Item LC 116 x NBS"])
    t.append(["ITEM LC 116", "DESCRIÇÃO ITEM", "NBS", "DESCRIÇÃO NBS", "PS ONEROSA", "ADQ EXTERIOR", "INDOP",
              "LOCAL INCIDÊNCIA IBS", "CCLASSTRIB", "NOME CCLASSTRIB"])
    t.append(["01.07", "Suporte", "1.1501.30.00", "Suporte técnico", "S", "N", "100301", "Local do adquirente",
              "000001", "Tributação integral"])
    t.append([None, None, "1.1508.00.00", "Manutenção", "S", "N", "100301", "Local", "000001", "Integral"])
    arq_viii = tmp_path / "AnexoVIII.xlsx"
    wb2.save(arq_viii)

    r = importar_arquivo(str(arq_b))
    r = importar_arquivo(str(arq_viii), r)
    assert CodigoTributacaoNacional.objects.get(codigo="010701").item_lc116 == "01.07"
    assert CodigoServicoLC116.objects.filter(item="13.05").exists()
    assert CodigoNBS.objects.get(codigo="115013000").codigo_mascarado == "1.1501.30.00"
    assert not CodigoNBS.objects.filter(codigo__startswith="115").exclude(codigo__regex=r"^\d{9}$").exists()
    # forward-fill da célula mesclada
    assert CorrelacaoNBS.objects.filter(item_lc116="01.07").count() == 2
    assert CorrelacaoNBS.objects.get(nbs="115013000").c_class_trib == "000001"
    assert r.correlacoes == 2


def test_preco_500_cartoes_com_variacoes(catalogo):
    v = {x.opcao: x.pk for x in catalogo.variacoes.all()}
    calc = calcular_preco(catalogo, 500, [v["Couché 300g"], v["Verniz localizado"],
                                          v["Faca especial (cantos arredondados)"]])
    assert calc["unitario"] == Decimal("0.8200")
    assert calc["total"] == Decimal("410.00")
    assert calcular_preco(catalogo, 1000)["unitario"] == Decimal("0.3800")
    assert calcular_preco(catalogo, 150)["unitario"] == Decimal("0.9000")


def test_sugestoes_fiscais(catalogo):
    s = sugestoes_fiscais("01.03")
    assert len(s["nbs"]) == 3 and len(s["codigo_tributacao_nacional"]) == 1


def test_item_pendencias_fiscais(catalogo):
    assert catalogo.pendencias_fiscais == ["NBS não informado (obrigatório desde 01/2026)."]
    catalogo.natureza = "PRODUTO"
    assert "NF-e" in catalogo.pendencias_fiscais[0]
    assert ItemCatalogo.objects.get(codigo_interno="TI-SUP").apto_nfse


def test_fluxo_orcamento_aprovacao_publica(cli, catalogo, pessoa):
    from django.test import Client

    client = Client()
    o = Orcamento.objects.create(cliente=pessoa, validade=timezone.localdate() + timedelta(days=10))
    assert o.numero.startswith(f"ORC-{timezone.localdate().year}-")
    v = {x.opcao: x.pk for x in catalogo.variacoes.all()}
    r = cli.post(reverse("comercial:item_adicionar", args=[o.pk]),
                 {"item_catalogo": catalogo.pk, "quantidade": "500",
                  "variacoes": [v["Couché 300g"], v["Verniz localizado"], v["Faca especial (cantos arredondados)"]]})
    assert r.status_code == 302
    o.refresh_from_db()
    assert o.total == Decimal("410.00")
    cli.post(reverse("comercial:enviar", args=[o.pk]), {"email": "compras@cliente.com"})
    o.refresh_from_db()
    assert o.status == "ENVIADO" and o.pdf_id and mail.outbox[-1].attachments

    # cliente abre no celular (sem login)
    url = reverse("comercial_publico:publico", args=[o.token_publico])
    r = client.get(url, HTTP_USER_AGENT="Mozilla/5.0 (iPhone)")
    assert r.status_code == 200 and "410,00" in r.content.decode()
    o.refresh_from_db()
    assert o.status == "VISUALIZADO"
    client.post(url, {"acao": "aprovar", "nome": "Maria da Silva"}, REMOTE_ADDR="200.1.2.3")
    o.refresh_from_db()
    assert o.status == "APROVADO" and o.aprovado_por_nome == "Maria da Silva" and o.ip_aprovacao == "200.1.2.3"

    r = cli.post(reverse("comercial:converter", args=[o.pk]))
    o.refresh_from_db()
    assert o.status == "CONVERTIDO" and o.venda.total == Decimal("410.00") and o.venda.itens.count() == 1


def test_orcamento_rascunho_nao_publico(client, empresa, pessoa):
    o = Orcamento.objects.create(cliente=pessoa, validade=timezone.localdate())
    assert client.get(reverse("comercial_publico:publico", args=[o.token_publico])).status_code == 404


def test_expiracao_e_followup(empresa, pessoa, admin_user):
    o = Orcamento.objects.create(cliente=pessoa, validade=timezone.localdate() - timedelta(days=1), status="ENVIADO")
    o2 = Orcamento.objects.create(cliente=pessoa, validade=timezone.localdate() + timedelta(days=9), status="ENVIADO",
                                  enviado_em=timezone.now() - timedelta(days=4))
    assert services.expirar_orcamentos() == 1
    o.refresh_from_db()
    assert o.status == "EXPIRADO"
    assert services.followup() == 1 and services.followup() == 0
    o2.refresh_from_db()
    assert o2.followup_em


def test_paginas_fase3(cli, catalogo, pessoa):
    o = Orcamento.objects.create(cliente=pessoa, validade=timezone.localdate() + timedelta(days=5))
    for url in [reverse("catalogo:item_lista"), reverse("catalogo:item_novo"), reverse("catalogo:tabelas") + "?q=suporte",
                reverse("catalogo:item_detalhe", args=[catalogo.pk]), reverse("comercial:orcamento_lista"),
                reverse("comercial:orcamento_novo"), reverse("comercial:orcamento_detalhe", args=[o.pk]),
                reverse("comercial:venda_lista"), reverse("comercial:preco") + f"?item_catalogo={catalogo.pk}&quantidade=500",
                reverse("comercial:variacoes") + f"?item_catalogo={catalogo.pk}",
                reverse("catalogo:sugestoes") + "?valor=01.07"]:
        assert cli.get(url).status_code == 200, url
