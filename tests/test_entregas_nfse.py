import io
import json
import smtplib
import zipfile
from datetime import timedelta
from decimal import Decimal

import pytest
from django.core import mail
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.base import ContentFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from apps.certidoes.models import Certidao, TipoCertidao
from apps.financeiro.services.custos import PARAMETROS, nova_planilha, salvar_planilha
from apps.fiscal.models import EntregaNota
from apps.fiscal.services import emissao, entregas, pdf_nfse
from apps.operacao.relatorios import novo_relatorio, salvar_relatorio
from tests.test_fiscal_integrado import arquivo_a1 as fixture_a1
from tests.test_fiscal_integrado import preparado as fixture_preparado
from tests.test_fiscal_integrado import resposta_nfse

arquivo_a1 = fixture_a1
preparado = fixture_preparado


def pdf_minimo():
    buf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture
def nota_autorizada(preparado):
    nota, config = preparado
    tentativa, a1 = emissao.preparar_tentativa(nota, config)
    return emissao.registrar_resultado(tentativa, 200, resposta_nfse(tentativa, a1), consulta=True)


@pytest.fixture
def pdf_rapido(monkeypatch):
    import apps.financeiro.services.custos as custos
    import apps.operacao.relatorios as relatorios

    for modulo in (pdf_nfse, entregas, custos, relatorios):
        monkeypatch.setattr(modulo, "renderizar_pdf", lambda *args, **kwargs: pdf_minimo())


@pytest.fixture
def pacote_contrato(nota_autorizada, contrato, admin_user, pdf_rapido):
    nota = nota_autorizada
    contrato.cliente = nota.tomador
    contrato.save()
    nota.contrato = contrato
    nota.save()
    tipo = TipoCertidao.objects.create(empresa=nota.empresa, nome="Certidao de teste", orgao_emissor="Teste", esfera="FEDERAL")
    contrato.certidoes_exigidas.add(tipo)
    hoje = timezone.localdate()
    cert = Certidao.objects.create(empresa=nota.empresa, tipo=tipo, data_emissao=hoje-timedelta(days=5),
                                   data_validade=hoje+timedelta(days=30), arquivo=ContentFile(pdf_minimo(), name="certidao.pdf"))
    rel = novo_relatorio(contrato, nota.competencia)
    rel = salvar_relatorio(rel, {"introducao": "Teste", "conteudo": "Atividades sinteticas de teste.", "consideracoes_finais": ""}, admin_user, True)
    plan = nova_planilha(contrato, nota.competencia)
    dados = dict.fromkeys(PARAMETROS, Decimal(0)) | {"tributos": Decimal(6), "observacoes": "Massa sintetica"}
    linhas = [{"grupo": "LICENCAS", "descricao": "=texto cadastrado", "unidade": "un", "quantidade": Decimal(2),
               "valor_unitario": Decimal(120), "rateio_pct": Decimal(50), "periodicidade": "ANUAL", "fonte": "Documento sintetico"}]
    plan = salvar_planilha(plan, dados, linhas, admin_user, True)
    return nota, cert, rel, plan


def preparar(nota):
    return entregas.preparar_entrega(nota, "destino@example.com; destino@example.com", "Nota de teste", "Mensagem de teste")


def test_pdf_real_uma_pagina_arquivado_e_cache(nota_autorizada):
    anexo = pdf_nfse.garantir_pdf(nota_autorizada)
    paginas = PdfReader(io.BytesIO(anexo.ler())).pages
    assert len(paginas) == 1
    texto = paginas[0].extract_text()
    assert "DANFSe v2.0" in texto and "TOMADOR SINTETICO" in texto
    assert "11111111111111111111111111111111111111111111111111" in texto
    assert pdf_nfse.garantir_pdf(nota_autorizada).pk == anexo.pk


def test_envio_avulso_pdf_xml_simulado_e_sem_repeticao(nota_autorizada, pdf_rapido):
    envio = preparar(nota_autorizada)
    assert envio.destinatarios == ["destino@example.com"]
    assert envio.status == "PREPARADA"
    envio = entregas.enviar_entrega(envio)
    assert envio.status == "SIMULADA" and envio.enviado_em
    assert len(mail.outbox) == 1
    assert {a.mimetype for a in mail.outbox[0].attachments} == {"application/pdf", "application/xml"}
    with pytest.raises(ValidationError):
        entregas.enviar_entrega(envio)
    assert len(mail.outbox) == 1


def test_pacote_contratual_completo_e_formula_bdi(pacote_contrato):
    nota, cert, rel, plan = pacote_contrato
    envio = preparar(nota)
    with zipfile.ZipFile(io.BytesIO(envio.pacote.ler())) as pacote:
        nomes = pacote.namelist()
        assert "kit_certidoes.zip" in nomes and rel.pdf.nome in nomes and plan.xlsx.nome in nomes and plan.pdf.nome in nomes
        manifesto = json.loads(pacote.read("manifesto.json"))
        assert manifesto["certidoes"][0]["id"] == str(cert.pk)
        with zipfile.ZipFile(io.BytesIO(pacote.read("kit_certidoes.zip"))) as kit:
            assert len(kit.namelist()) == 2 and "00_indice.pdf" in kit.namelist()
    wb = load_workbook(io.BytesIO(plan.xlsx.ler()))
    assert wb.active["B4"].value == "'=texto cadastrado"
    assert wb.active["H4"].value == "=ROUND(D4*E4*F4/12,2)"
    valores = load_workbook(io.BytesIO(plan.xlsx.ler()), data_only=True).active
    assert valores["H4"].value == 10 and valores["H6"].value == 10
    assert valores["H16"].value == 10.64
    assert plan.itens.get().valor_mensal == Decimal("10.00")
    assert abs(plan.bdi - Decimal(6)/94) < Decimal("0.00000001")
    entregas.enviar_entrega(envio)
    assert len(mail.outbox[0].attachments) == 1 and mail.outbox[0].attachments[0].mimetype == "application/zip"


def test_certidao_expirada_apos_preparo_bloqueia_envio(pacote_contrato):
    nota, cert, _, _ = pacote_contrato
    envio = preparar(nota)
    cert.data_validade = timezone.localdate()-timedelta(days=1)
    cert.save()
    with pytest.raises(ValidationError):
        entregas.enviar_entrega(envio)
    envio.refresh_from_db()
    assert envio.status == "PREPARADA"
    assert not getattr(mail, "outbox", [])


def test_nova_versao_aprovada_exige_novo_pacote(pacote_contrato, admin_user):
    nota, _, rel, _ = pacote_contrato
    envio = preparar(nota)
    with pytest.raises(ValidationError):
        salvar_relatorio(rel, {"introducao": "", "conteudo": "Alterado", "consideracoes_finais": ""}, admin_user)
    novo = novo_relatorio(nota.contrato, nota.competencia)
    assert novo.versao == 2
    salvar_relatorio(novo, {"introducao": "", "conteudo": "Nova versao sintetica", "consideracoes_finais": ""}, admin_user, True)
    with pytest.raises(ValidationError):
        entregas.enviar_entrega(envio)


def test_nova_certidao_exigida_apos_preparo_bloqueia(pacote_contrato):
    nota, cert, _, _ = pacote_contrato
    envio = preparar(nota)
    tipo = TipoCertidao.objects.create(empresa=nota.empresa, nome="Outra exigencia", orgao_emissor="Teste", esfera="OUTRA")
    Certidao.objects.create(empresa=nota.empresa, tipo=tipo, data_emissao=cert.data_emissao, data_validade=cert.data_validade,
                            arquivo=ContentFile(pdf_minimo(), name="nova.pdf"))
    nota.contrato.certidoes_exigidas.add(tipo)
    with pytest.raises(ValidationError, match="outra certid"):
        entregas.enviar_entrega(envio)


@pytest.mark.parametrize("alteracao", ["POSITIVA", "SEM_ARQUIVO", "VENCIDA"])
def test_certidoes_inadequadas_nao_entram_no_kit(pacote_contrato, alteracao):
    nota, cert, _, _ = pacote_contrato
    if alteracao == "POSITIVA":
        cert.situacao = "POSITIVA"
    elif alteracao == "SEM_ARQUIVO":
        cert.arquivo = ""
    else:
        cert.data_validade = timezone.localdate()-timedelta(days=1)
    cert.save()
    with pytest.raises(ValidationError):
        preparar(nota)
    assert not EntregaNota.objects.exists()


def test_contrato_sem_documentos_nao_monta_pacote(nota_autorizada, contrato, pdf_rapido):
    contrato.cliente = nota_autorizada.tomador
    contrato.save()
    nota_autorizada.contrato = contrato
    nota_autorizada.save()
    with pytest.raises(ValidationError, match="atividades"):
        preparar(nota_autorizada)
    assert not EntregaNota.objects.exists()


def test_xml_adulterado_bloqueia_pdf(nota_autorizada, pdf_rapido):
    arquivo = nota_autorizada.xml_autorizado
    arquivo.arquivo.save("alterado.xml", ContentFile(b"<alterado/>"))
    with pytest.raises(ValidationError, match="alterado"):
        pdf_nfse.garantir_pdf(nota_autorizada)


def test_sla_exigido_bloqueia_pacote_incompleto(pacote_contrato):
    nota, _, _, _ = pacote_contrato
    nota.contrato.exige_relatorio_sla = True
    nota.contrato.save()
    with pytest.raises(ValidationError, match="SLA"):
        preparar(nota)


def test_homologacao_nao_envia_email_real(nota_autorizada):
    nota_autorizada.ambiente = 2
    nota_autorizada.save()
    with pytest.raises(ValidationError, match="produ"):
        preparar(nota_autorizada)


def test_relatorio_so_inclui_fatos_e_horas_da_competencia(contrato, admin_user):
    from datetime import date

    from apps.operacao.models import Apontamento, Tarefa

    tarefa = Tarefa.objects.create(empresa=contrato.empresa, contrato=contrato, responsavel=admin_user,
          titulo="Suporte da competencia", resumo_para_relatorio="Entrega sintetica", data_abertura=date(2026, 9, 1))
    for data, horas in ((date(2026, 9, 30), Decimal(8)), (date(2026, 10, 1), Decimal("1.50"))):
        Apontamento.objects.create(empresa=contrato.empresa, tarefa=tarefa, usuario=admin_user, data=data, horas=horas, descricao="Teste")
    Tarefa.objects.create(empresa=contrato.empresa, contrato=contrato, responsavel=admin_user, titulo="Sem execucao registrada")
    rel = novo_relatorio(contrato, date(2026, 10, 1))
    assert len(rel.tarefas_snapshot) == 1 and rel.tarefas_snapshot[0]["horas"] == "1.50"
    assert "Sem execucao registrada" not in rel.conteudo


def test_documentos_e_downloads_respeitam_papeis_e_empresa(pacote_contrato, cli, leitor, django_user_model):
    from django.contrib.auth.models import Group

    from apps.core.models import Empresa

    nota, _, rel, plan = pacote_contrato
    with pytest.raises(PermissionDenied):
        salvar_relatorio(rel, {}, leitor)
    with pytest.raises(PermissionDenied):
        salvar_planilha(plan, {}, [], leitor)
    tecnico = django_user_model.objects.create_user(email="tecnico-documentos@example.com", empresa=nota.empresa)
    tecnico.groups.add(Group.objects.get(name="Operação"))
    cli.force_login(tecnico)
    assert cli.get(reverse("core:anexo_baixar", args=[rel.pdf_id])).status_code == 403
    nota.contrato.responsaveis.add(tecnico)
    assert cli.get(reverse("core:anexo_baixar", args=[rel.pdf_id])).status_code == 200
    assert cli.get(reverse("core:anexo_baixar", args=[plan.pdf_id])).status_code == 403
    outra = Empresa.objects.create(cnpj="11222333000181", razao_social="Outra empresa sintetica")
    usuario = django_user_model.objects.create_user(email="outra-documentos@example.com", empresa=outra)
    usuario.groups.add(Group.objects.get(name="Administrador"))
    cli.force_login(usuario)
    assert cli.get(reverse("fiscal:pdf", args=[nota.pk])).status_code == 404
    assert cli.get(reverse("fiscal:documentos_contrato", args=[nota.contrato_id])).status_code == 404


def test_smtp_sem_confirmacao_nao_reenvia(nota_autorizada, pdf_rapido, monkeypatch, settings):
    settings.EMAIL_BACKEND = "django.core.mail.backends.smtp.EmailBackend"
    envio = preparar(nota_autorizada)
    def erro(*args, **kwargs):
        raise smtplib.SMTPServerDisconnected("teste sintetico")
    monkeypatch.setattr(entregas.EmailMessage, "send", erro)
    with pytest.raises(ValidationError, match="sem confirma"):
        entregas.enviar_entrega(envio)
    envio.refresh_from_db()
    assert envio.status == "INCERTA"
    with pytest.raises(ValidationError, match="iniciado"):
        entregas.enviar_entrega(envio)


def test_smtp_mock_aceite_nunca_afirma_entrega(nota_autorizada, pdf_rapido, monkeypatch, settings):
    settings.EMAIL_BACKEND = "django.core.mail.backends.smtp.EmailBackend"
    monkeypatch.setattr(entregas.EmailMessage, "send", lambda *args, **kwargs: 1)
    envio = entregas.enviar_entrega(preparar(nota_autorizada))
    assert envio.status == "ENVIADA" and "servidor" in envio.get_status_display()


def test_permissoes_e_telas(cli, client, leitor, nota_autorizada, contrato, pdf_rapido):
    from apps.financeiro.models import PlanilhaCustos
    from apps.operacao.models import RelatorioAtividades

    resposta_pdf = cli.get(reverse("fiscal:pdf", args=[nota_autorizada.pk]))
    assert resposta_pdf.status_code == 200 and resposta_pdf["Content-Type"] == "application/pdf"
    assert "object-src 'none'" in resposta_pdf["Content-Security-Policy"]
    resposta_html = cli.get(reverse("fiscal:entregas", args=[nota_autorizada.pk]))
    assert resposta_html.status_code == 200 and "object-src 'none'" in resposta_html["Content-Security-Policy"]
    url = reverse("fiscal:documentos_contrato", args=[contrato.pk]) + "?competencia=2026-10-01"
    assert cli.get(url).status_code == 200
    assert cli.post(url, {"acao": "novo_relatorio"}).status_code == 302
    assert cli.post(url, {"acao": "nova_planilha"}).status_code == 302
    assert cli.get(url).status_code == 200
    rel = RelatorioAtividades.objects.get(contrato=contrato)
    assert cli.post(url, {"acao": "aprovar_relatorio", "versao_id": str(rel.pk), "introducao": "", "conteudo": "Teste de formulario", "consideracoes_finais": ""}).status_code == 302
    rel.refresh_from_db()
    assert rel.status == "APROVADO" and rel.pdf_id
    plan = PlanilhaCustos.objects.get(contrato=contrato)
    dados = {campo: "0" for campo in PARAMETROS} | {"observacoes": "Teste", "acao": "aprovar_custos", "versao_id": str(plan.pk),
              "custos-TOTAL_FORMS": "1", "custos-INITIAL_FORMS": "0", "custos-MIN_NUM_FORMS": "0", "custos-MAX_NUM_FORMS": "100",
              "custos-0-grupo": "LICENCAS", "custos-0-descricao": "Teste", "custos-0-unidade": "un", "custos-0-quantidade": "1",
              "custos-0-valor_unitario": "10", "custos-0-rateio_pct": "100", "custos-0-periodicidade": "MENSAL", "custos-0-fonte": "Teste"}
    resposta = cli.post(url, dados)
    assert resposta.status_code == 302
    plan.refresh_from_db()
    assert plan.status == "APROVADO" and plan.pdf_id and plan.xlsx_id
    resposta = cli.post(reverse("fiscal:entregas", args=[nota_autorizada.pk]), {"destinatarios": "destino@example.com", "assunto": "Teste", "mensagem": "Teste"})
    assert resposta.status_code == 302
    assert cli.get(resposta.url).status_code == 200
    client.force_login(leitor)
    assert client.get(reverse("fiscal:pdf", args=[nota_autorizada.pk])).status_code == 200
    assert client.post(reverse("fiscal:entregas", args=[nota_autorizada.pk]), {}).status_code == 403
    assert client.post(url, {"acao": "novo_relatorio"}).status_code == 403
