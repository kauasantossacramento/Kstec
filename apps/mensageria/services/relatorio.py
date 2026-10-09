"""Resumo financeiro em XLSX para envio pelo WhatsApp (comando RELATORIO)."""

import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.models import Anexo
from apps.faturamento.services import previsao
from apps.financeiro.models import Lancamento
from apps.fiscal.models import NotaFiscal

AZUL = PatternFill("solid", fgColor="0016E1")
BRL = '"R$" #,##0.00'


def _texto_seguro(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _aba(wb, titulo, cabecalho, linhas, moedas=()):
    ws = wb.create_sheet(titulo)
    ws.append(cabecalho)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), AZUL, Alignment(vertical="center")
    for linha in linhas:
        ws.append([_texto_seguro(v) for v in linha])
    for col in moedas:
        for cel in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cel[0].number_format = BRL
    for i, _ in enumerate(cabecalho, 1):
        ws.column_dimensions[get_column_letter(i)].width = 18 if i > 1 else 44
    ws.freeze_panes = "A2"
    return ws


def gerar(empresa, hoje=None):
    hoje = hoje or timezone.localdate()
    dados = previsao.mes(empresa, hoje)
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    resumo.append([f"Resumo financeiro · {empresa} · gerado em {timezone.localtime():%d/%m/%Y %H:%M}"])
    resumo["A1"].font = Font(bold=True, size=13)
    linhas = [("Previsto no mês", dados["total"]), ("Recebido no mês", dados["totais"]["RECEBIDO"]),
              ("Lançado em aberto", dados["totais"]["EM_ABERTO"]), ("Programado (faturamento)", dados["totais"]["PROGRAMADO"]),
              ("Estimado (contratos)", dados["totais"]["ESTIMADO"]),
              ("Em atraso de meses anteriores", dados["valor_atrasados_anteriores"])]
    for rotulo, valor in linhas:
        resumo.append([rotulo, valor])
        resumo.cell(resumo.max_row, 2).number_format = BRL
    resumo.column_dimensions["A"].width = 38
    resumo.column_dimensions["B"].width = 18
    _aba(wb, "Previsão do mês", ["Descrição", "Data", "Origem", "Cliente", "Valor", "Ajustado"],
         [(i.titulo, i.data, i.rotulo_origem, i.cliente, i.valor, "sim" if i.ajuste else "")
          for i in dados["itens"] if not i.excluido], moedas=(5,))
    abertos = Lancamento.objects.filter(empresa=empresa, tipo="RECEITA", status__in=["PREVISTO", "PENDENTE", "ATRASADO"]) \
        .select_related("pessoa").order_by("data_vencimento")
    _aba(wb, "Em aberto", ["Descrição", "Vencimento", "Cliente", "Valor", "Situação"],
         [(x.descricao, x.data_vencimento, str(x.pessoa or ""), x.valor, x.get_status_display()) for x in abertos],
         moedas=(4,))
    notas = NotaFiscal.objects.filter(empresa=empresa, competencia__year=hoje.year).select_related("tomador")
    _aba(wb, "Notas do ano", ["Tomador", "Competência", "Número", "Valor", "Situação"],
         [(str(n.tomador), n.competencia, n.numero_nfse, n.valor_servicos, n.get_status_display()) for n in notas],
         moedas=(4,))
    for ws in wb.worksheets[1:]:
        for linha in ws.iter_rows(min_row=2):
            for cel in linha:
                if hasattr(cel.value, "year") and not isinstance(cel.value, Decimal):
                    cel.number_format = "DD/MM/YYYY"
    saida = io.BytesIO()
    wb.save(saida)
    return Anexo.criar(None, f"resumo-financeiro-{hoje:%Y-%m-%d}.xlsx", saida.getvalue(),
                       descricao="Resumo financeiro solicitado pelo WhatsApp", empresa=empresa)
