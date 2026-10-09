"""Importação das tabelas fiscais de referência (Fase 3, seção 1.2).

- `ANEXO_B-NBS2-LISTA_SERVICO_NACIONAL.xlsx`: abas `LISTA.SERV.NAC.` (LC 116 / cTribNac) e `LISTA.NBS_v2.0` (NBS).
- `AnexoVIII-CorrelacaoItemNBSIndOpCClassTrib_IBSCBS.xlsx`: aba `tabela geral` (Item LC 116 → NBS → IndOp →
  cClassTrib), com *forward-fill* da coluna do item LC 116, que vem mesclada.

Os layouts dos anexos mudam entre versões; por isso o importador localiza a linha de cabeçalho e as colunas por
palavras-chave, e reconhece os códigos por padrão (01.07, 01.07.01, 1.1502.20.00) em vez de posição fixa.
"""

import re
import unicodedata
from dataclasses import dataclass, field

from django.db import transaction
from openpyxl import load_workbook

from .models import CodigoNBS, CodigoServicoLC116, CodigoTributacaoNacional, CorrelacaoNBS

RE_ITEM = re.compile(r"^(\d{1,2})\.(\d{2})$")
RE_DESDOBRO = re.compile(r"^(\d{1,2})\.(\d{2})\.(\d{2})$")
RE_CTRIB6 = re.compile(r"^\d{6}$")
RE_NBS = re.compile(r"^\d\.\d{4}\.\d{2}\.\d{2}$")
RE_NBS9 = re.compile(r"^\d{9}$")


@dataclass
class Resultado:
    lc116: int = 0
    ctribnac: int = 0
    nbs: int = 0
    correlacoes: int = 0
    avisos: list[str] = field(default_factory=list)

    def __str__(self):
        return (f"LC 116: {self.lc116} · cTribNac: {self.ctribnac} · NBS: {self.nbs} · "
                f"correlações: {self.correlacoes}")


def _norm(v) -> str:
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", s).strip()


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _aba(wb, *trechos):
    for nome in wb.sheetnames:
        n = _norm(nome).replace(" ", "")
        if any(t.replace(" ", "") in n for t in trechos):
            return wb[nome]
    return None


def _descricao(celulas: list[str], usados: set[int]) -> str:
    candidatos = [(len(c), c) for i, c in enumerate(celulas) if i not in usados and c and not re.match(r"^[\d.]+$", c)]
    return max(candidatos)[1] if candidatos else ""


def importar_lista_servicos(ws, r: Resultado):
    descricoes_item = {}
    colunas = None
    total_antes = r.ctribnac
    for linha in ws.iter_rows(values_only=True):
        cel = [_txt(c) for c in linha]
        normalizadas = [_norm(c) for c in cel]
        if "CODIGO DE TRIBUTACAO NACIONAL" in normalizadas and "SUBITEM" in normalizadas:
            colunas = {nome: normalizadas.index(nome) for nome in
                       ["CODIGO DE TRIBUTACAO NACIONAL", "ITEM", "SUBITEM", "DESDOBRO NACIONAL", "DESCRICAO"]}
            CodigoTributacaoNacional.objects.update(vigente=False)
            continue
        if colunas is not None:
            bruto = cel[colunas["CODIGO DE TRIBUTACAO NACIONAL"]]
            it, sub = cel[colunas["ITEM"]], cel[colunas["SUBITEM"]]
            desc = cel[colunas["DESCRICAO"]]
            if not it.isdigit() or not sub.isdigit() or int(sub) == 0:
                continue  # Cabeçalhos de grupo (01.00) não são serviços tributáveis.
            item = f"{int(it):02d}.{int(sub):02d}"
            descricoes_item.setdefault(item, desc)
            if not bruto:
                # Linha de subitem da LC 116: o desdobro 00 é agrupador, não cTribNac.
                CodigoServicoLC116.objects.update_or_create(item=item, defaults={"descricao": desc})
                r.lc116 += 1
                continue
            cod = bruto.replace(".", "").zfill(6)
            if not RE_CTRIB6.fullmatch(cod):
                raise ValueError(f"Código nacional inválido no anexo: {bruto}")
            CodigoTributacaoNacional.objects.update_or_create(codigo=cod, defaults={
                "item": cod[:2], "subitem": cod[2:4], "desdobro": cod[4:6], "descricao": desc, "vigente": True})
            r.ctribnac += 1
            continue
        # Formato A: código completo numa célula ("01.07.01", "010701" ou "01.07")
        for i, c in enumerate(cel):
            m3 = RE_DESDOBRO.match(c)
            m6 = RE_CTRIB6.match(c) if not m3 else None
            m2 = RE_ITEM.match(c) if not (m3 or m6) else None
            if m3 or m6:
                cod = f"{int(m3.group(1)):02d}{m3.group(2)}{m3.group(3)}" if m3 else c
                desc = _descricao(cel, {i})
                CodigoTributacaoNacional.objects.update_or_create(
                    codigo=cod, defaults={"item": cod[:2], "subitem": cod[2:4], "desdobro": cod[4:6], "descricao": desc})
                r.ctribnac += 1
                descricoes_item.setdefault(f"{cod[:2]}.{cod[2:4]}", desc)
                break
            if m2:
                item = f"{int(m2.group(1)):02d}.{m2.group(2)}"
                desc = _descricao(cel, {i})
                CodigoServicoLC116.objects.update_or_create(item=item, defaults={"descricao": desc})
                descricoes_item[item] = desc
                r.lc116 += 1
                break
        else:
            # Formato B: colunas separadas item | subitem | desdobro | descrição
            nums = [(i, c) for i, c in enumerate(cel) if re.fullmatch(r"\d{1,2}", c)]
            if len(nums) >= 3:
                (i1, it), (i2, sub), (i3, des) = nums[:3]
                cod = f"{int(it):02d}{int(sub):02d}{int(des):02d}"
                desc = _descricao(cel, {i1, i2, i3})
                CodigoTributacaoNacional.objects.update_or_create(
                    codigo=cod, defaults={"item": cod[:2], "subitem": cod[2:4], "desdobro": cod[4:6], "descricao": desc})
                r.ctribnac += 1
                descricoes_item.setdefault(f"{cod[:2]}.{cod[2:4]}", desc)
    if colunas is not None and r.ctribnac == total_antes:
        raise ValueError("Anexo sem códigos nacionais; importação cancelada para preservar a lista anterior.")
    # Itens da LC 116 que só apareceram via desdobro
    for item, desc in descricoes_item.items():
        _, criado = CodigoServicoLC116.objects.get_or_create(item=item, defaults={"descricao": desc})
        r.lc116 += criado


def importar_nbs(ws, r: Resultado):
    for linha in ws.iter_rows(values_only=True):
        cel = [_txt(c) for c in linha]
        for i, c in enumerate(cel):
            c2 = c.replace(" ", "")
            if RE_NBS.match(c2) or RE_NBS9.match(c2):
                codigo = c2.replace(".", "")
                CodigoNBS.objects.update_or_create(codigo=codigo, defaults={
                    "codigo_mascarado": CodigoNBS.mascarar(codigo), "descricao": _descricao(cel, {i})})
                r.nbs += 1
                break


COLUNAS_CORRELACAO = {
    "item_lc116": ["ITEM LC", "ITEM DA LC", "LC 116", "LC116", "ITEM"],
    "nbs": ["NBS"],
    "descricao_nbs": ["DESCRICAO NBS", "DESCRICAO DA NBS", "DESCRICAO"],
    "ps_onerosa": ["ONEROSA"],
    "adq_exterior": ["EXTERIOR"],
    "ind_op": ["INDOP", "IND OP", "INDICADOR DA OPERACAO", "INDICADOR"],
    "local_incidencia_ibs": ["LOCAL"],
    "c_class_trib": ["CCLASSTRIB", "C CLASS TRIB", "CLASSTRIB", "CLASSIFICACAO TRIBUTARIA"],
    "nome_c_class_trib": ["NOME CCLASSTRIB", "NOME DA CCLASSTRIB", "NOME CLASS", "DESCRICAO CCLASSTRIB"],
}


def _mapear_cabecalho(cel: list[str]) -> dict:
    norm = [_norm(c) for c in cel]
    mapa = {}
    # ordem importa: campos mais específicos primeiro
    for campo in ["nome_c_class_trib", "c_class_trib", "descricao_nbs", "ps_onerosa", "adq_exterior", "ind_op",
                  "local_incidencia_ibs", "nbs", "item_lc116"]:
        for chave in COLUNAS_CORRELACAO[campo]:
            for i, n in enumerate(norm):
                if i in mapa.values():
                    continue
                if chave in n.replace(".", "") and not (campo == "nbs" and "DESCRI" in n):
                    mapa[campo] = i
                    break
            if campo in mapa:
                break
    return mapa


def importar_correlacao(ws, r: Resultado):
    mapa = None
    item_atual = ""
    CorrelacaoNBS.objects.all().delete()
    lote = []
    for linha in ws.iter_rows(values_only=True):
        cel = [_txt(c) for c in linha]
        if mapa is None:
            m = _mapear_cabecalho(cel)
            if "nbs" in m and "item_lc116" in m:
                mapa = m
            continue
        bruto_item = cel[mapa["item_lc116"]] if mapa["item_lc116"] < len(cel) else ""
        m_item = re.match(r"^(\d{1,2})[.,](\d{1,2})", bruto_item)
        if m_item:
            item_atual = f"{int(m_item.group(1)):02d}.{int(m_item.group(2)):02d}"  # forward-fill (célula mesclada)
        nbs = re.sub(r"\D", "", cel[mapa["nbs"]] if mapa["nbs"] < len(cel) else "")
        if not item_atual or len(nbs) != 9:
            continue
        dados = {campo: (cel[i] if i < len(cel) else "")[:300] for campo, i in mapa.items()
                 if campo not in ("item_lc116", "nbs")}
        lote.append(CorrelacaoNBS(item_lc116=item_atual, nbs=nbs, **dados))
    if not lote:
        raise ValueError("Nenhuma correlação válida encontrada; tabela anterior preservada.")
    CorrelacaoNBS.objects.bulk_create(lote, batch_size=1000)
    r.correlacoes += len(lote)
    if mapa is None:
        r.avisos.append("Cabeçalho da 'tabela geral' não encontrado (colunas Item LC 116 e NBS).")


@transaction.atomic
def importar_arquivo(caminho_ou_arquivo, r: Resultado | None = None) -> Resultado:
    r = r or Resultado()
    wb = load_workbook(caminho_ou_arquivo, read_only=True, data_only=True)
    reconhecido = False
    ws = _aba(wb, "LISTA.SERV.NAC", "LISTASERVNAC", "SERV.NAC")
    if ws is not None:
        importar_lista_servicos(ws, r)
        reconhecido = True
    ws = _aba(wb, "LISTA.NBS", "NBS_V2", "NBS")
    if ws is not None and not _aba(wb, "TABELA GERAL"):
        importar_nbs(ws, r)
        reconhecido = True
    ws = _aba(wb, "TABELA GERAL")
    if ws is not None:
        importar_correlacao(ws, r)
        reconhecido = True
    if not reconhecido:
        r.avisos.append(f"Nenhuma aba reconhecida. Abas: {', '.join(wb.sheetnames)}")
    wb.close()
    return r
