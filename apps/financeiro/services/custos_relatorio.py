"""Relatório de custos: por contrato (custo direto + rateio dos custos globais), por centro e por categoria.

Regime de competência (data_competencia), lançamentos não cancelados. Custos de centros sem contrato e marcados
como "ratear" são distribuídos entre os contratos proporcionalmente à receita do período; sem receita, em partes
iguais entre os contratos com movimento ou vigentes.
"""

import io
from decimal import ROUND_HALF_UP, Decimal

from django.db.models import Q, Sum
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.contratos.models import Contrato

from ..models import CentroCusto, Lancamento

ZERO = Decimal("0")
CENT = Decimal("0.01")


def _soma(qs):
    return qs.aggregate(v=Sum("valor"))["v"] or ZERO


def apurar(empresa, inicio, fim, contrato=None):
    base = Lancamento.objects.filter(empresa=empresa, data_competencia__range=(inicio, fim)).exclude(status="CANCELADO")
    despesas = base.filter(tipo="DESPESA")
    receitas = base.filter(tipo="RECEITA")
    centros_contrato = CentroCusto.objects.filter(empresa=empresa, contrato__isnull=False).select_related("contrato__cliente")
    globais = CentroCusto.objects.filter(empresa=empresa, contrato__isnull=True)
    total_global = _soma(despesas.filter(centro_custo__in=globais.filter(ratear=True)))
    nao_rateado = _soma(despesas.filter(centro_custo__in=globais.filter(ratear=False)))
    linhas = []
    for centro in centros_contrato:
        c = centro.contrato
        receita = _soma(receitas.filter(centro_custo=centro))
        direto = _soma(despesas.filter(centro_custo=centro))
        vigente = c.status == "VIGENTE" and c.vigencia_inicio <= fim and c.vigencia_fim_atual >= inicio
        if receita or direto or vigente:
            linhas.append({"contrato": c, "centro": centro, "receita": receita, "direto": direto})
    total_receita = sum((x["receita"] for x in linhas), ZERO)
    for x in linhas:
        if total_receita:
            x["rateio"] = (total_global * x["receita"] / total_receita).quantize(CENT, ROUND_HALF_UP)
        else:
            x["rateio"] = (total_global / len(linhas)).quantize(CENT, ROUND_HALF_UP) if linhas else ZERO
        x["custo_total"] = x["direto"] + x["rateio"]
        x["resultado"] = x["receita"] - x["custo_total"]
        x["margem"] = (x["resultado"] / x["receita"] * 100).quantize(Decimal("0.1")) if x["receita"] else None
    if contrato is not None:
        linhas = [x for x in linhas if x["contrato"].pk == contrato.pk]
    centros = []
    for centro in CentroCusto.objects.filter(empresa=empresa).select_related("contrato"):
        valor = _soma(despesas.filter(centro_custo=centro))
        if valor:
            centros.append({"centro": centro, "valor": valor, "pago": _soma(despesas.filter(centro_custo=centro, status="PAGO")),
                            "global": centro.contrato_id is None, "rateado": centro.contrato_id is None and centro.ratear})
    filtro_cat = despesas if contrato is None else despesas.filter(
        Q(centro_custo__contrato=contrato) | Q(centro_custo__contrato__isnull=True))
    categorias = list(filtro_cat.values("categoria__nome").annotate(valor=Sum("valor")).order_by("-valor"))
    itens = despesas.select_related("centro_custo", "categoria", "pessoa").order_by("data_competencia")
    if contrato is not None:
        itens = itens.filter(Q(centro_custo__contrato=contrato) | Q(centro_custo__contrato__isnull=True, centro_custo__ratear=True))
    return {"inicio": inicio, "fim": fim, "contrato": contrato, "linhas": linhas, "centros": centros,
            "categorias": categorias, "itens": itens,
            "total_custos": _soma(despesas), "total_global": total_global, "nao_rateado": nao_rateado,
            "total_direto": sum((x["direto"] for x in linhas), ZERO), "total_receita": total_receita,
            "resultado": total_receita - _soma(despesas),
            "criterio": "proporcional à receita" if total_receita else "partes iguais (sem receita no período)"}


def xlsx(dados):
    wb = Workbook()
    azul = PatternFill("solid", fgColor="0016E1")
    brl = '"R$" #,##0.00'

    def aba(ws, cab, linhas, moedas):
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), azul
        for linha in linhas:
            ws.append([("'" + v) if isinstance(v, str) and v[:1] in "=+-@" else v for v in linha])
        for col in moedas:
            for (cel,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cel.number_format = brl
        ws.column_dimensions["A"].width = 42
        for letra in "BCDEFGH":
            ws.column_dimensions[letra].width = 16

    ws = wb.active
    ws.title = "Por contrato"
    aba(ws, ["Contrato", "Cliente", "Receita", "Custo direto", "Rateio global", "Custo total", "Resultado", "Margem %"],
        [(x["contrato"].numero, str(x["contrato"].cliente), x["receita"], x["direto"], x["rateio"], x["custo_total"],
          x["resultado"], float(x["margem"]) if x["margem"] is not None else None) for x in dados["linhas"]], (3, 4, 5, 6, 7))
    aba(wb.create_sheet("Por centro de custo"), ["Centro de custo", "Tipo", "Custo", "Pago"],
        [(x["centro"].nome, "Global (rateado)" if x["rateado"] else "Global" if x["global"] else "Contrato", x["valor"], x["pago"])
         for x in dados["centros"]], (3, 4))
    aba(wb.create_sheet("Por categoria"), ["Categoria", "Custo"],
        [(x["categoria__nome"], x["valor"]) for x in dados["categorias"]], (2,))
    aba(wb.create_sheet("Lançamentos"), ["Descrição", "Competência", "Centro", "Categoria", "Fornecedor", "Valor", "Situação"],
        [(i.descricao, i.data_competencia, i.centro_custo.nome, i.categoria.nome, str(i.pessoa or ""), i.valor,
          i.get_status_display()) for i in dados["itens"]], (6,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def contratos_para_filtro(empresa):
    return Contrato.objects.filter(empresa=empresa).select_related("cliente")
